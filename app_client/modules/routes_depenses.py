# -*- coding: utf-8 -*-
"""Routes de gestion des dépenses : analyse, fournisseurs, CRUD, scan facture, export."""

import os
import io
import json
from datetime import date, datetime, timedelta

from flask import render_template, request, redirect, url_for, jsonify, make_response, send_file
from flask_login import login_required, current_user

from modules.config import app
from modules.db import get_db
from modules.helpers import _rid, _mois_courant, _mois_label, _get_mois_list, _stats_mois, _depenses_mois, _seuil_rentabilite
from modules import airtable_client as at
from modules import airtable_sync as sync


# ==============================================================================
#  Analyse Dépenses
# ==============================================================================

@app.route("/gestion/analyse-depenses")
def analyse_depenses():
    return render_template("base.html", page="analyse_depenses")


@app.route("/gestion/fournisseurs", methods=["GET", "POST"])
def gestion_fournisseurs():
    db = get_db()
    if request.method == "POST":
        action = request.form.get("action", "create")
        if action == "create":
            nom = request.form.get("nom", "").strip()
            if nom:
                sync.write_fournisseur(db, nom, request.form.get("ftype", "Autre"),
                                       telephone=request.form.get("telephone", ""),
                                       email=request.form.get("email", ""),
                                       nom_commercial=request.form.get("nom_commercial", ""),
                                       telephone_commercial=request.form.get("telephone_commercial", ""),
                                       email_commercial=request.form.get("email_commercial", ""),
                                       site_web=request.form.get("site_web", ""),
                                       adresse=request.form.get("adresse", ""),
                                       notes=request.form.get("notes", ""))
        elif action == "update":
            fid = int(request.form["id"])
            sync.update_fournisseur(db, fid, request.form.get("nom", "").strip(),
                                    request.form.get("ftype", "Autre"),
                                    telephone=request.form.get("telephone", ""),
                                    email=request.form.get("email", ""),
                                    nom_commercial=request.form.get("nom_commercial", ""),
                                    telephone_commercial=request.form.get("telephone_commercial", ""),
                                    email_commercial=request.form.get("email_commercial", ""),
                                    site_web=request.form.get("site_web", ""),
                                    adresse=request.form.get("adresse", ""),
                                    notes=request.form.get("notes", ""))
        elif action == "delete":
            fid = int(request.form.get("id", 0) or 0)
            sync.delete_fournisseur(db, fid)
        db.commit()
        return redirect(url_for("gestion_fournisseurs"))

    fournisseurs = db.execute("SELECT * FROM fournisseurs WHERE restaurant_id=? ORDER BY nom", (_rid(),)).fetchall()
    fournisseurs_data = []
    mois_courant = _mois_courant()
    for f in fournisseurs:
        dep_mois = db.execute("SELECT SUM(montant) as t FROM depenses_variables WHERE restaurant_id=? AND description=? AND mois=?",
                              (_rid(), f["nom"], mois_courant)).fetchone()
        total_mois = round(dep_mois["t"] or 0, 2)
        dep_all = db.execute("SELECT SUM(montant) as t FROM depenses_variables WHERE restaurant_id=? AND description=?",
                             (_rid(), f["nom"],)).fetchone()
        total_all = round(dep_all["t"] or 0, 2)
        hist = db.execute("SELECT date, montant, categorie FROM depenses_variables WHERE restaurant_id=? AND description=? ORDER BY date DESC LIMIT 20",
                          (_rid(), f["nom"],)).fetchall()
        fournisseurs_data.append({
            **dict(f),
            "total_mois": total_mois,
            "total_all": total_all,
            "historique": [dict(h) for h in hist],
        })
    return render_template("base.html", page="fournisseurs_page", fournisseurs_list=fournisseurs_data)


@app.route("/gestion/analyse-depenses/data")
def analyse_depenses_data():
    db = get_db()
    periode = request.args.get("periode", "mois")
    today = date.today()

    # Calculer les bornes de la periode
    if periode == "mois":
        date_from = today.strftime("%Y-%m-01")
        pm = today.month - 1
        py = today.year
        if pm <= 0:
            pm = 12
            py -= 1
        prev_from = f"{py:04d}-{pm:02d}-01"
        prev_to = (date(today.year, today.month, 1) - timedelta(days=1)).isoformat()
    elif periode == "3mois":
        d = today - timedelta(days=90)
        date_from = d.isoformat()
        prev_from = (d - timedelta(days=90)).isoformat()
        prev_to = (d - timedelta(days=1)).isoformat()
    elif periode == "6mois":
        d = today - timedelta(days=180)
        date_from = d.isoformat()
        prev_from = (d - timedelta(days=180)).isoformat()
        prev_to = (d - timedelta(days=1)).isoformat()
    elif periode == "annee":
        date_from = f"{today.year}-01-01"
        prev_from = f"{today.year-1}-01-01"
        prev_to = f"{today.year-1}-12-31"
    else:  # alltime
        date_from = "2000-01-01"
        prev_from = None
        prev_to = None
    date_to = today.isoformat()

    def _mois_from_date(d):
        return d[:7] if d and len(d) >= 7 else ""

    # Compute months in period (needed for fixes)
    mois_list_periode = set()
    d_iter = date.fromisoformat(date_from) if date_from > "2000-01-01" else date(today.year, today.month, 1)
    while d_iter <= today:
        mois_list_periode.add(d_iter.strftime("%Y-%m"))
        if d_iter.month == 12:
            d_iter = d_iter.replace(year=d_iter.year + 1, month=1)
        else:
            d_iter = d_iter.replace(month=d_iter.month + 1)

    # --- PAR FOURNISSEUR ---
    known_fourns = [r["nom"] for r in db.execute("SELECT nom FROM fournisseurs WHERE restaurant_id=? ORDER BY nom", (_rid(),)).fetchall()]
    mois_list_sorted = sorted(mois_list_periode)
    if known_fourns and mois_list_sorted:
        mois_placeholders = ",".join("?" * len(mois_list_sorted))
        fournisseurs_rows = db.execute("""
            SELECT dv.description as fournisseur, dv.categorie, SUM(dv.montant) as total, COUNT(*) as nb
            FROM depenses_variables dv
            WHERE dv.mois IN ({}) AND dv.description IN ({})
            GROUP BY dv.description ORDER BY total DESC
        """.format(mois_placeholders, ",".join("?" * len(known_fourns))),
            (*mois_list_sorted, *known_fourns)).fetchall()
    else:
        fournisseurs_rows = []

    rec_fixes = db.execute("SELECT description, categorie, montant FROM depenses_fixes WHERE restaurant_id=? AND recurrente=1", (_rid(),)).fetchall()
    fixes_as_fourns = {}
    nb_mois = len(mois_list_periode) if mois_list_periode else 1
    for r in rec_fixes:
        key = r["description"] or r["categorie"]
        if key not in fixes_as_fourns:
            fixes_as_fourns[key] = {"categorie": r["categorie"], "total": 0, "nb": 0}
        fixes_as_fourns[key]["total"] += r["montant"] * nb_mois
        fixes_as_fourns[key]["nb"] += nb_mois

    all_fourns = []
    total_var = 0
    for r in fournisseurs_rows:
        t = round(r["total"], 2)
        total_var += t
        all_fourns.append({"categorie": r["fournisseur"], "type": r["categorie"], "total": t,
                           "nb": r["nb"], "moyenne": round(t / r["nb"], 2) if r["nb"] > 0 else 0, "pct": 0})
    all_fourns.sort(key=lambda x: -x["total"])
    for f in all_fourns:
        f["pct"] = round(f["total"] / total_var * 100, 1) if total_var > 0 else 0
    fournisseurs = all_fourns

    # Repartition par categorie
    type_totals = {}
    for f in fournisseurs:
        t = f["type"]
        type_totals[t] = type_totals.get(t, 0) + f["total"]
    type_repartition = [{"type": k, "total": round(v, 2)} for k, v in sorted(type_totals.items(), key=lambda x: -x[1])]

    # --- PAR CATEGORIE ---
    total_all = 0
    total_variables = 0
    total_fixes = 0
    agg_food = 0
    agg_bev = 0
    agg_personnel = 0
    for mv in sorted(mois_list_periode):
        dep_m = _depenses_mois(db, mv)
        total_all += dep_m["total"]
        total_variables += dep_m["total_variables"]
        total_fixes += dep_m["total_fixes"]
        agg_food += dep_m["food"]
        agg_bev += dep_m["beverage"]
        agg_personnel += dep_m["personnel"]
    total_all = round(total_all, 2)
    total_variables = round(total_variables, 2)
    total_fixes = round(total_fixes, 2)
    agg_food = round(agg_food, 2)
    agg_bev = round(agg_bev, 2)

    ca_periode = 0
    for mv in mois_list_sorted:
        s = _stats_mois(db, mv)
        ca_periode += s["ca_total"]
    ca_periode = round(ca_periode, 2)

    resultat_net = round((ca_periode - total_all) / ca_periode * 100, 1) if ca_periode > 0 else 0

    # Periode precedente
    prev_mois_list = set()
    if prev_from and prev_to:
        try:
            d_iter2 = date.fromisoformat(prev_from)
            prev_end = date.fromisoformat(prev_to)
            while d_iter2 <= prev_end:
                prev_mois_list.add(d_iter2.strftime("%Y-%m"))
                if d_iter2.month == 12:
                    d_iter2 = d_iter2.replace(year=d_iter2.year + 1, month=1)
                else:
                    d_iter2 = d_iter2.replace(month=d_iter2.month + 1)
        except Exception:
            pass

    prev_all = 0
    for mv2 in sorted(prev_mois_list):
        dep_p = _depenses_mois(db, mv2)
        prev_all += dep_p["total"]
    prev_all = round(prev_all, 2)
    evo_pct = round((total_all - prev_all) / prev_all * 100, 1) if prev_all > 0 else 0

    categories = [
        {"nom": "Cout matieres", "total": agg_food, "pct_ca": round(agg_food / ca_periode * 100, 1) if ca_periode > 0 else 0},
        {"nom": "Cout boissons", "total": agg_bev, "pct_ca": round(agg_bev / ca_periode * 100, 1) if ca_periode > 0 else 0},
        {"nom": "Personnel", "total": round(agg_personnel, 2), "pct_ca": round(agg_personnel / ca_periode * 100, 1) if ca_periode > 0 else 0},
        {"nom": "Charges fixes", "total": total_fixes, "pct_ca": round(total_fixes / ca_periode * 100, 1) if ca_periode > 0 else 0},
    ]
    categories = [c for c in categories if c["total"] > 0]

    all_var = db.execute("SELECT SUM(montant) as t FROM depenses_variables WHERE restaurant_id=?", (_rid(),)).fetchone()
    all_fix = db.execute("SELECT SUM(montant) as t FROM depenses_fixes WHERE restaurant_id=?", (_rid(),)).fetchone()
    all_time_total = round((all_var["t"] or 0) + (all_fix["t"] or 0), 2)

    top5 = db.execute("SELECT categorie, SUM(montant) as total FROM depenses_variables WHERE restaurant_id=? GROUP BY categorie ORDER BY total DESC LIMIT 5", (_rid(),)).fetchall()
    top5_list = [{"nom": r["categorie"], "total": round(r["total"], 2)} for r in top5]

    # Par annee
    annees_data = []
    years = set()
    for r in db.execute("SELECT DISTINCT substr(date,1,4) as y FROM depenses_variables WHERE restaurant_id=?", (_rid(),)).fetchall():
        years.add(r["y"])
    for r in db.execute("SELECT DISTINCT substr(mois,1,4) as y FROM depenses_fixes WHERE restaurant_id=?", (_rid(),)).fetchall():
        years.add(r["y"])
    for y in sorted(years):
        y_dep_total = 0
        y_var_total = 0
        y_fix_total = 0
        for m in range(1, 13):
            mv = f"{y}-{m:02d}"
            dep_ym = _depenses_mois(db, mv)
            y_dep_total += dep_ym["total"]
            y_var_total += dep_ym["total_variables"]
            y_fix_total += dep_ym["total_fixes"]
        dep_y = round(y_dep_total, 2)
        yca = db.execute("SELECT SUM(ca) as t FROM ca_jour WHERE restaurant_id=? AND date LIKE ?", (_rid(), y + "%",)).fetchone()
        ca_y = round(yca["t"], 2) if yca["t"] else 0
        annees_data.append({"annee": y, "dépenses": dep_y, "var": round(y_var_total, 2), "fix": round(y_fix_total, 2), "ca": ca_y, "ratio": round(dep_y / ca_y * 100, 1) if ca_y > 0 else 0})

    # Insight
    insight = None
    if top5_list:
        top = top5_list[0]
        pct_top = round(top["total"] / all_time_total * 100, 1) if all_time_total > 0 else 0
        insight = {"nom": top["nom"], "total": top["total"], "pct": pct_top}

    return jsonify({
        "fournisseurs": fournisseurs,
        "type_repartition": type_repartition,
        "categories": categories,
        "total_all": total_all,
        "ca_periode": ca_periode,
        "evo_pct": evo_pct,
        "resultat_net": resultat_net,
        "alltime_total": all_time_total,
        "top5": top5_list,
        "annees": annees_data,
        "insight": insight,
        "is_demo": False,
    })


# ==============================================================================
#  Dépenses CRUD
# ==============================================================================

@app.route("/depenses", methods=["GET", "POST"])
def depenses():
    mois = request.args.get("mois", _mois_courant())
    db = get_db()

    if request.method == "POST":
        form_type = request.form.get("form_type")
        if form_type == "fixe":
            cat = request.form.get("categorie", "")
            desc = request.form.get("description", "")
            montant = float(request.form.get("montant", 0))
            recurrente = 1 if request.form.get("recurrente") else 0
            sync.write_depense_fixe(db, cat, desc, montant, mois, recurrente)
        elif form_type == "variable":
            d = request.form.get("date", date.today().isoformat())
            cat = request.form.get("categorie", "")
            desc = request.form.get("description", "")
            montant = float(request.form.get("montant", 0) or 0)
            recurrente = 1 if request.form.get("recurrente") else 0
            frequence = request.form.get("frequence", "mensuelle") if recurrente else ""
            sync.write_depense_variable(db, d, cat, desc, montant, mois)
            if recurrente:
                last_id = db.execute("SELECT id FROM depenses_variables WHERE restaurant_id=? ORDER BY id DESC LIMIT 1", (_rid(),)).fetchone()
                if last_id:
                    db.execute("UPDATE depenses_variables SET recurrente=1, frequence=? WHERE id=?",
                               (frequence, last_id["id"]))
                    db.commit()
        elif form_type == "fournisseur":
            nom = request.form.get("nom", "").strip()
            ftype = request.form.get("ftype", "Autre")
            if nom:
                sync.write_fournisseur(db, nom, ftype)
        return redirect(url_for("depenses", mois=mois))

    # Direct fixes for this month
    fixes_direct = db.execute("SELECT * FROM depenses_fixes WHERE restaurant_id=? AND mois = ? ORDER BY categorie", (_rid(), mois,)).fetchall()
    recurrentes = db.execute("""SELECT * FROM depenses_fixes
        WHERE restaurant_id=? AND recurrente = 1 AND mois < ? ORDER BY categorie""", (_rid(), mois,)).fetchall()

    fixes_ids_direct = {f["id"] for f in fixes_direct}
    direct_keys = {(f["categorie"], f["description"]) for f in fixes_direct}

    fixes_list = []
    for f in fixes_direct:
        entry = dict(f)
        entry["is_recurrente"] = bool(f["recurrente"]) if "recurrente" in f.keys() else False
        entry["is_inherited"] = False
        fixes_list.append(entry)

    for r in recurrentes:
        key = (r["categorie"], r["description"])
        if key not in direct_keys:
            desact = (r["desactivee_mois"] or "")
            if mois in desact.split(","):
                continue
            entry = dict(r)
            entry["is_recurrente"] = True
            entry["is_inherited"] = True
            entry["source_mois"] = r["mois"]
            fixes_list.append(entry)

    variables = db.execute("SELECT * FROM depenses_variables WHERE restaurant_id=? AND mois = ? ORDER BY date DESC", (_rid(), mois,)).fetchall()
    total_fixes = sum(f["montant"] for f in fixes_list)
    total_variables = sum(r["montant"] for r in variables)
    fournisseurs = db.execute("SELECT * FROM fournisseurs WHERE restaurant_id=? ORDER BY nom", (_rid(),)).fetchall()

    _FOOD_CATS = {'Food', 'Matières premières', 'Emballages', 'Autre achats',
                  'Hygiène & Entretien', 'Entretien', 'Énergie', 'Réparations',
                  'Reparation', 'Autre charges opé', 'Fournitures', 'Marketing', 'Livraison'}
    _BEV_CATS = {'Boissons'}
    _PERSONNEL_CATS = {'Salaires', 'Extras & Intérim', 'Extras / Interim',
                       'Formation', 'Autre personnel', 'Charges sociales'}
    food = round(sum(v["montant"] for v in variables if v["categorie"] in _FOOD_CATS), 2)
    beverage = round(sum(v["montant"] for v in variables if v["categorie"] in _BEV_CATS), 2)
    personnel = round(sum(v["montant"] for v in variables if v["categorie"] in _PERSONNEL_CATS), 2)
    stats = _stats_mois(db, mois)
    ca = stats["ca_total"]
    total_dep = round(total_fixes + total_variables, 2)
    recap = {
        "food": food, "beverage": beverage, "personnel": personnel,
        "total_fixes": round(total_fixes, 2), "total_dep": total_dep, "ca": ca,
        "food_pct": round(food / ca * 100, 1) if ca > 0 else 0,
        "bev_pct": round(beverage / ca * 100, 1) if ca > 0 else 0,
        "personnel_pct": round(personnel / ca * 100, 1) if ca > 0 else 0,
        "fixes_pct": round(total_fixes / ca * 100, 1) if ca > 0 else 0,
        "total_pct": round(total_dep / ca * 100, 1) if ca > 0 else 0,
        "resultat": round(ca - total_dep, 2),
        "resultat_pct": round((ca - total_dep) / ca * 100, 1) if ca > 0 else 0,
    }
    return render_template("base.html", page="depenses",
        mois=mois, mois_label=_mois_label(mois), mois_list=_get_mois_list(),
        fixes=fixes_list, variables=variables, fournisseurs=fournisseurs,
        total_fixes=round(total_fixes, 2), total_variables=round(total_variables, 2),
        recap=recap)


@app.route("/depenses/scan", methods=["POST"])
def depenses_scan():
    """Analyse une photo de facture via Claude Vision et retourne les donnees extraites."""
    import base64
    rid = getattr(current_user, "restaurant_id", "")
    api_key = os.environ.get("ANTHROPIC_API_KEY", "")
    if not api_key:
        return jsonify({"ok": False, "error": "Clé API Anthropic non configurée. Contactez le support.",
                        "error_type": "config"}), 200

    photo = request.files.get("photo")
    if not photo or not photo.filename:
        return jsonify({"ok": False, "error": "Aucune photo reçue.", "error_type": "input"}), 200

    img_bytes = photo.read()
    print(f"[SCAN] Photo recue: {photo.filename}, type={photo.content_type}, size={len(img_bytes)} bytes")
    if len(img_bytes) > 20 * 1024 * 1024:
        return jsonify({"ok": False, "error": "Image trop volumineuse (max 20 Mo).", "error_type": "input"}), 200
    img_b64 = base64.b64encode(img_bytes).decode("utf-8")
    content_type = photo.content_type or "image/jpeg"
    supported_types = {"image/jpeg", "image/png", "image/webp", "image/gif"}
    if content_type not in supported_types:
        print(f"[SCAN] Content-type non supporté '{content_type}', force image/jpeg")
        content_type = "image/jpeg"

    # Upload vers R2
    photo_key = ""
    photo_url = ""
    from modules.r2_client import upload_file, is_configured as r2_ok, get_file_url
    if r2_ok():
        try:
            import uuid
            from io import BytesIO
            ext = photo.filename.rsplit(".", 1)[-1] if "." in photo.filename else "jpg"
            now = datetime.now()
            fname = f"{now.strftime('%Y%m%d_%H%M%S')}_{uuid.uuid4().hex[:6]}.{ext}"
            folder = f"factures/{rid}/{now.strftime('%Y')}/{now.strftime('%m')}"
            fobj = BytesIO(img_bytes)
            fobj.content_type = content_type
            photo_key = f"{folder}/{fname}"
            upload_file(photo_key, fobj)
            photo_url = get_file_url(photo_key)
            print(f"[SCAN] Uploaded to R2: {photo_key}")
        except Exception as e:
            print(f"[SCAN] R2 upload error: {e}")

    # Appel Claude Vision
    try:
        import requests as _rq
        payload = {
            "model": "claude-sonnet-4-20250514",
            "max_tokens": 4000,
            "messages": [{
                "role": "user",
                "content": [
                    {
                        "type": "text",
                        "text": (
                            "Tu es un assistant expert en factures de restaurant. "
                            "Extrais les informations suivantes de cette facture au format JSON UNIQUEMENT, sans commentaire. "
                            "Utilise ces clés exactes : fournisseur, date, montant_ht, tva, montant_ttc, numero_facture, produits. "
                            "Le montant est en euros. "
                            "Si tu ne trouves pas une information, mets une chaîne vide."
                        )
                    },
                    {
                        "type": "image",
                        "source": {
                            "type": "base64",
                            "media_type": content_type,
                            "data": img_b64
                        }
                    }
                ]
            }]
        }
        headers = {
            "x-api-key": api_key,
            "anthropic-version": "2023-06-01",
            "content-type": "application/json"
        }
        resp = _rq.post("https://api.anthropic.com/v1/messages", headers=headers, json=payload, timeout=60)
        resp.raise_for_status()
        result = resp.json()
        text = result["content"][0]["text"]
        print(f"[SCAN] Claude response: {text[:200]}...")
        # Extract JSON from response
        json_start = text.find("{")
        json_end = text.rfind("}")
        if json_start >= 0 and json_end > json_start:
            json_str = text[json_start:json_end + 1]
            facture = json.loads(json_str)
        else:
            return jsonify({"ok": False, "error": "Réponse Claude invalide.", "error_type": "parse"}), 200
    except Exception as e:
        return jsonify({"ok": False, "error": f"Erreur analyse: {e}", "error_type": "api"}), 200

    # Normaliser fournisseur : chercher dans la table fournisseurs
    fourn_raw = facture.get("fournisseur", "")
    fourn_match = None
    fournisseurs_list = []
    if fourn_raw:
        # Chercher une correspondance
        fourn_upper = fourn_raw.strip().upper()
        all_fourns = db.execute("SELECT id, nom FROM fournisseurs WHERE restaurant_id=? ORDER BY nom", (_rid(),)).fetchall()
        for f in all_fourns:
            fn = f["nom"].strip().upper()
            # 1) Exact match
            if fn == fourn_upper:
                fourn_match = f["nom"]
                break
            # 2) Contains match (e.g. "Metro Saint-Malo" -> "Metro")
            if fn in fourn_upper or fourn_upper in fn:
                fourn_match = f["nom"]
                break
            # 3) First-word fuzzy match (e.g. "METRO CASH & CARRY" -> "Metro")
            if len(fn) >= 3 and fn.split()[0] == fourn_upper.split()[0]:
                fourn_match = f["nom"]
                break
        fournisseurs_list = [{"id": f["id"], "nom": f["nom"], "selected": f["nom"] == fourn_match} for f in all_fourns]
        if not fourn_match:
            fournisseurs_list = [{"id": f["id"], "nom": f["nom"], "selected": False} for f in all_fourns]

    return jsonify({
        "ok": True,
        "facture": facture,
        "fournisseur_match": fourn_match,
        "fournisseur_raw": facture.get("fournisseur", ""),
        "fournisseurs": fournisseurs_list,
        "photo_key": photo_key,
        "photo_url": photo_url,
    })


@app.route("/depenses/scan-save", methods=["POST"])
def depenses_scan_save():
    """Sauvegarde une facture scannée dans SQLite + Airtable."""
    data = request.get_json(silent=True) or {}
    db = get_db()
    mois = data.get("mois") or _mois_courant()
    date_str = data.get("date", date.today().isoformat())
    cat = data.get("categorie", "Food")
    fourn = data.get("fournisseur", "")
    montant = float(data.get("montant", 0))
    num_facture = data.get("numero_facture", "")
    montant_ht = data.get("montant_ht", "")
    tva = data.get("tva", "")
    produits = data.get("produits", "")
    commentaire = data.get("commentaire", "")
    photo_key = data.get("photo_key", "")

    desc_parts = []
    if fourn:
        desc_parts.append(fourn)
    if num_facture:
        desc_parts.append(f"Fact. {num_facture}")
    if produits:
        desc_parts.append(produits)
    if commentaire:
        desc_parts.append(commentaire)
    description = " — ".join(desc_parts) if desc_parts else fourn

    mois_val = date_str[:7] if date_str and len(date_str) >= 7 else mois

    try:
        sync.write_depense_variable(db, date_str, cat, description, montant, mois_val)
        return jsonify({"ok": True})
    except Exception as e:
        print(f"[SCAN-SAVE] Error: {e}")
        return jsonify({"ok": False, "error": str(e)}), 200


@app.route("/depenses/fournisseur-quick", methods=["POST"])
def depenses_fournisseur_quick():
    """Création rapide d'un fournisseur depuis le scan facture."""
    data = request.get_json(silent=True) or {}
    nom = data.get("nom", "").strip()
    ftype = data.get("type", "Autre")
    if not nom:
        return jsonify({"ok": False, "error": "Nom requis"}), 200
    db = get_db()
    try:
        sync.write_fournisseur(db, nom, ftype)
        new_f = db.execute("SELECT id, nom, type FROM fournisseurs WHERE restaurant_id=? AND nom=?", (_rid(), nom,)).fetchone()
        return jsonify({"ok": True, "fournisseur": {"id": new_f["id"], "nom": new_f["nom"], "type": new_f["type"]} if new_f else {"nom": nom}})
    except Exception as e:
        return jsonify({"ok": False, "error": str(e)}), 200


@app.route("/depenses/export")
def depenses_export():
    db = get_db()
    fmt = request.args.get("format", "csv")
    periode = request.args.get("periode", "mois")
    today = date.today()

    if periode == "mois":
        date_from = today.strftime("%Y-%m-01")
        date_to = today.isoformat()
        mois_from = _mois_courant()
    elif periode == "trimestre":
        d = today - timedelta(days=90)
        date_from = d.isoformat()
        date_to = today.isoformat()
        mois_from = d.strftime("%Y-%m")
    elif periode == "annee":
        date_from = f"{today.year}-01-01"
        date_to = today.isoformat()
        mois_from = f"{today.year}-01"
    else:
        date_from = "2000-01-01"
        date_to = today.isoformat()
        mois_from = "2000-01"

    rows = []
    mois_to = today.strftime("%Y-%m")

    def _months_between(start_ym, end_ym):
        sy, sm = int(start_ym[:4]), int(start_ym[5:7])
        ey, em = int(end_ym[:4]), int(end_ym[5:7])
        out = []
        y, m = sy, sm
        while (y, m) <= (ey, em):
            out.append(f"{y:04d}-{m:02d}")
            m += 1
            if m > 12:
                m = 1
                y += 1
        return out

    target_months = _months_between(mois_from, mois_to)

    for r in db.execute("SELECT * FROM depenses_variables WHERE restaurant_id=? AND (mois >= ? AND mois <= ? OR mois = '' OR mois IS NULL) ORDER BY date", (_rid(), mois_from, mois_to)).fetchall():
        rows.append({"date": r["date"] or "", "categorie": r["categorie"] or "", "fournisseur": r["description"] or "", "montant": r["montant"] or 0, "type": "Variable"})

    fixes = db.execute("SELECT * FROM depenses_fixes WHERE restaurant_id=? ORDER BY mois", (_rid(),)).fetchall()
    for r in fixes:
        mois_creation = r["mois"] or ""
        if not mois_creation:
            continue
        recurrente = bool(r["recurrente"]) if "recurrente" in r.keys() else False
        desactivee = (r["desactivee_mois"] or "").split(",") if "desactivee_mois" in r.keys() else []
        for tm in target_months:
            if mois_creation > tm:
                continue
            if not recurrente and mois_creation != tm:
                continue
            if tm in desactivee:
                continue
            rows.append({"date": tm, "categorie": r["categorie"] or "", "fournisseur": r["description"] or "", "montant": r["montant"] or 0, "type": "Fixe"})
    rows.sort(key=lambda x: x["date"])

    def _fmt_date(s):
        s = (s or "").strip()
        if not s:
            return ""
        if len(s) >= 10 and s[4] == "-" and s[7] == "-":
            return f"{s[8:10]}/{s[5:7]}/{s[:4]}"
        return s

    if fmt == "csv":
        import csv
        output = io.StringIO()
        writer = csv.writer(output, delimiter=";")
        writer.writerow(["Date", "Catégorie", "Fournisseur", "Montant", "Type"])
        for r in rows:
            writer.writerow([_fmt_date(r["date"]), r["categorie"], r["fournisseur"], f"{r['montant']:.2f}".replace(".", ","), r["type"]])
        csv_data = output.getvalue()
        output.close()
        resp = make_response(csv_data.encode("utf-8-sig"))
        resp.headers["Content-Type"] = "text/csv; charset=utf-8"
        resp.headers["Content-Disposition"] = f"attachment; filename=depenses-{periode}.csv"
        return resp

    elif fmt == "xlsx":
        try:
            import openpyxl
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Depenses"
            ws.append(["Date", "Catégorie", "Fournisseur", "Montant", "Type"])
            for r in rows:
                ws.append([_fmt_date(r["date"]), r["categorie"], r["fournisseur"], r["montant"], r["type"]])
            buf = io.BytesIO()
            wb.save(buf)
            buf.seek(0)
            resp = make_response(buf.getvalue())
            resp.headers["Content-Type"] = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            resp.headers["Content-Disposition"] = f"attachment; filename=depenses-{periode}.xlsx"
            return resp
        except ImportError:
            return "openpyxl non installé", 500

    elif fmt == "pdf":
        try:
            from reportlab.lib.pagesizes import A4
            from reportlab.lib import colors
            from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
            from reportlab.lib.styles import getSampleStyleSheet
            from reportlab.lib.units import mm
        except ImportError:
            return "reportlab non installé", 500

        buf = io.BytesIO()
        doc = SimpleDocTemplate(buf, pagesize=A4, topMargin=15 * mm, bottomMargin=15 * mm,
                                leftMargin=10 * mm, rightMargin=10 * mm)
        styles = getSampleStyleSheet()
        elems = []
        elems.append(Paragraph(f"Export dépenses - {periode}", styles["Title"]))
        elems.append(Spacer(1, 5 * mm))

        data_table = [["Date", "Catégorie", "Fournisseur", "Montant", "Type"]]
        for r in rows:
            data_table.append([_fmt_date(r["date"]), r["categorie"], r["fournisseur"], f"{r['montant']:.2f}", r["type"]])

        t = Table(data_table, colWidths=[30 * mm, 35 * mm, 45 * mm, 25 * mm, 20 * mm])
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2D6A4A")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, 0), 9),
            ("FONTSIZE", (0, 1), (-1, -1), 8),
            ("ALIGN", (3, 0), (3, -1), "RIGHT"),
            ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#E4DDD3")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F7F4EF")]),
        ]))
        elems.append(t)
        doc.build(elems)
        buf.seek(0)
        resp = make_response(buf.getvalue())
        resp.headers["Content-Type"] = "application/pdf"
        resp.headers["Content-Disposition"] = f"attachment; filename=depenses-{periode}.pdf"
        return resp

    return "Format non supporte", 400


@app.route("/depenses/delete/<table>/<int:id>")
@login_required
def delete_depense(table, id):
    mois = request.args.get("mois", _mois_courant())
    db = get_db()
    if table in ("fixe", "variable"):
        sync.delete_depense(db, table, id)
    elif table == "fournisseur":
        sync.delete_fournisseur(db, id)
    db.commit()
    return redirect(url_for("depenses", mois=mois))


@app.route("/depenses/disable-recurrente/<int:id>")
def disable_recurrente(id):
    """Desactive une depense recurrente pour un mois specifique."""
    mois = request.args.get("mois", _mois_courant())
    db = get_db()
    row = db.execute("SELECT desactivee_mois FROM depenses_fixes WHERE restaurant_id=? AND id=?", (_rid(), id,)).fetchone()
    if row:
        current = (row["desactivee_mois"] or "").split(",")
        current = [m for m in current if m]
        if mois not in current:
            current.append(mois)
        db.execute("UPDATE depenses_fixes SET desactivee_mois=? WHERE id=?", (",".join(current), id))
        db.commit()
    return redirect(url_for("depenses", mois=mois))


@app.route("/depenses/enable-recurrente/<int:id>")
def enable_recurrente(id):
    """Reactive une depense recurrente pour un mois specifique."""
    mois = request.args.get("mois", _mois_courant())
    db = get_db()
    row = db.execute("SELECT desactivee_mois FROM depenses_fixes WHERE restaurant_id=? AND id=?", (_rid(), id,)).fetchone()
    if row:
        current = (row["desactivee_mois"] or "").split(",")
        current = [m for m in current if m and m != mois]
        db.execute("UPDATE depenses_fixes SET desactivee_mois=? WHERE id=?", (",".join(current), id))
        db.commit()
    return redirect(url_for("depenses", mois=mois))
